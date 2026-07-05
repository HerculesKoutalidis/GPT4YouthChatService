"""
build_catalog.py
================
Builds / updates data/metadata/docs_catalog.xlsx — one row per PDF found in
data/raw_pdfs/. Prefills each row (best effort) from the two partner metadata
files, then leaves it for you to review and complete.

    python src/ingestion/build_catalog.py \
        --partners data/metadata/Partners_data_assets.xlsx \
        --cazalla  data/metadata/Cazala_sources_metadata.xlsx

Behaviour:
  * INCREMENTAL: existing rows in docs_catalog.xlsx are preserved (your edits
    are never overwritten). Only PDFs not yet in the catalog get new rows.
    So the "add PDFs one by one" workflow just works — drop new PDFs in
    data/raw_pdfs/, re-run this, fill the new rows.
  * Prefill priority: exact filename match (Partners) > fuzzy title match
    (Partners / Cazalla). Fuzzy matches are flagged 'fuzzy?' in
    metadata_source so you know which rows to double-check.

Partner filenames don't always match the real PDF names (e.g. Partners lists
'Compass_Manual_..._2023.txt' but the file is 'Compass 2023 ENG_final_WEB.pdf'),
which is exactly why prefill is best-effort + flagged.
"""

import os
import argparse
from difflib import SequenceMatcher

import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill

from catalog import CATALOG_COLUMNS, normalize

BASE_DIR = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
RAW_PDF_DIR = os.path.join(BASE_DIR, "data", "raw_pdfs")
DEFAULT_OUT = os.path.join(BASE_DIR, "data", "metadata", "docs_catalog.xlsx")

FUZZY_THRESHOLD = 0.60


# ---------------------------------------------------------------------------
# Load partner metadata into a common shape: list of {norm_keys, fields}
# ---------------------------------------------------------------------------
def _rows(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb.worksheets[0]
    data = list(ws.iter_rows(values_only=True))
    header = [str(c).strip() if c is not None else "" for c in data[0]]
    return header, data[1:]


def load_partners(path):
    """Infofront: Document_ID, File_name, Link, Title, Country, Summary, Tags,
    Who_it_helps, How_it_helps."""
    if not path or not os.path.exists(path):
        return []
    header, rows = _rows(path)
    idx = {h: i for i, h in enumerate(header)}

    def g(r, name):
        i = idx.get(name)
        return (r[i] if i is not None and i < len(r) and r[i] is not None else "")

    out = []
    for r in rows:
        title = str(g(r, "Title")).strip()
        fname = str(g(r, "File_name")).strip()
        if not (title or fname):
            continue
        out.append({
            "match_keys": {normalize(fname), normalize(os.path.splitext(fname)[0]), normalize(title)},
            "fields": {
                "title": title,
                "country": str(g(r, "Country")).strip(),
                "summary": str(g(r, "Summary")).strip(),
                "tags": str(g(r, "Tags")).strip(),
                "who_it_helps": str(g(r, "Who_it_helps")).strip(),
                "how_it_helps": str(g(r, "How_it_helps")).strip(),
                "source_link": str(g(r, "Link")).strip(),
            },
            "source": "partners",
        })
    return out


def load_cazalla(path):
    """Cazalla resource directory: Resource Title, Resource Type, Year, Language,
    Link / File, Short Description, Keywords, Target Audience, Main Use, ..."""
    if not path or not os.path.exists(path):
        return []
    header, rows = _rows(path)
    idx = {h: i for i, h in enumerate(header)}

    def g(r, name):
        i = idx.get(name)
        return (r[i] if i is not None and i < len(r) and r[i] is not None else "")

    out = []
    for r in rows:
        title = str(g(r, "Resource Title")).strip()
        if not title:
            continue
        year = str(g(r, "Year")).strip().replace(".0", "")
        out.append({
            "match_keys": {normalize(title)},
            "fields": {
                "title": title,
                "language": str(g(r, "Language")).strip().lower(),
                "doc_type": str(g(r, "Resource Type")).strip().lower(),
                "summary": str(g(r, "Short Description")).strip(),
                "tags": str(g(r, "Keywords")).strip(),
                "who_it_helps": str(g(r, "Target Audience")).strip(),
                "how_it_helps": str(g(r, "Main Use")).strip(),
                "source_link": str(g(r, "Link / File")).strip(),
                "year": year,
            },
            "source": "cazalla",
        })
    return out


# ---------------------------------------------------------------------------
# Matching a PDF filename to the best partner record
# ---------------------------------------------------------------------------
def best_match(stem_norm, records):
    exact = [r for r in records if stem_norm in r["match_keys"]]
    if exact:
        return exact[0], 1.0
    best, best_score = None, 0.0
    for r in records:
        score = max((SequenceMatcher(None, stem_norm, k).ratio() for k in r["match_keys"] if k), default=0.0)
        if score > best_score:
            best, best_score = r, score
    if best and best_score >= FUZZY_THRESHOLD:
        return best, best_score
    return None, 0.0


def detect_language(pdf_path):
    try:
        import fitz
        from langdetect import detect
        doc = fitz.open(pdf_path)
        text = "".join(doc[i].get_text() for i in range(min(5, doc.page_count)))
        doc.close()
        return detect(text[:3000]) if text.strip() else ""
    except Exception:
        return ""


def guess_doc_type(name):
    n = name.lower()
    for t in ("toolkit", "manual", "guide", "handbook", "report", "course"):
        if t in n:
            return "handbook" if t == "handbook" else t
    return ""


# ---------------------------------------------------------------------------
# Build / update
# ---------------------------------------------------------------------------
def build(args):
    if not os.path.isdir(RAW_PDF_DIR):
        raise SystemExit(f"PDF dir not found: {RAW_PDF_DIR}")

    from catalog import load_catalog
    existing = load_catalog(args.out)                      # preserve edits
    partners = load_partners(args.partners)
    cazalla = load_cazalla(args.cazalla)

    pdfs = sorted(f for f in os.listdir(RAW_PDF_DIR) if f.lower().endswith(".pdf"))
    added, kept = 0, 0
    catalog = dict(existing)

    for fn in pdfs:
        if fn in catalog:
            kept += 1
            continue
        stem = os.path.splitext(fn)[0]
        stem_norm = normalize(stem)
        row = {c: "" for c in CATALOG_COLUMNS}
        row["file_name"] = fn
        row["title"] = stem
        row["doc_type"] = guess_doc_type(fn)
        row["metadata_source"] = "MANUAL - fill me"

        # try Partners first, then Cazalla
        rec, score = best_match(stem_norm, partners)
        if not rec:
            rec, score = best_match(stem_norm, cazalla)
        if rec:
            for k, v in rec["fields"].items():
                if v:
                    row[k] = v
            row["metadata_source"] = rec["source"] if score >= 0.999 else f"{rec['source']} fuzzy? ({score:.2f})"

        if not row.get("language"):
            row["language"] = detect_language(os.path.join(RAW_PDF_DIR, fn))

        catalog[fn] = row
        added += 1

    write_xlsx(catalog, args.out)
    print(f"\nCatalog: {args.out}")
    print(f"  PDFs scanned : {len(pdfs)}")
    print(f"  rows kept    : {kept} (untouched)")
    print(f"  rows added   : {added}")
    print("  -> Review rows flagged 'MANUAL - fill me' or 'fuzzy?' before ingesting.")


def write_xlsx(catalog, out_path):
    os.makedirs(os.path.dirname(out_path), exist_ok=True)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "catalog"

    header_fill = PatternFill("solid", fgColor="DDDDDD")
    ws.append(CATALOG_COLUMNS)
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = header_fill
    ws.freeze_panes = "A2"

    for fn in sorted(catalog):
        row = catalog[fn]
        ws.append([row.get(c, "") for c in CATALOG_COLUMNS])

    # flag rows needing attention
    warn_fill = PatternFill("solid", fgColor="FFF2CC")
    ms_col = CATALOG_COLUMNS.index("metadata_source") + 1
    for r in range(2, ws.max_row + 1):
        val = str(ws.cell(r, ms_col).value or "")
        if "MANUAL" in val or "fuzzy" in val:
            ws.cell(r, ms_col).fill = warn_fill

    widths = {"file_name": 34, "title": 34, "summary": 50, "who_it_helps": 30,
              "how_it_helps": 30, "tags": 26, "source_link": 30, "metadata_source": 22}
    for i, c in enumerate(CATALOG_COLUMNS, start=1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = widths.get(c, 14)
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    # readme sheet
    rm = wb.create_sheet("_readme")
    notes = [
        ["Column", "Meaning"],
        ["file_name", "EXACT PDF filename in data/raw_pdfs/. This is the join key — do not change."],
        ["title", "Human-readable document title."],
        ["language", "ISO code (en, es, ...). Non-supported langs are skipped at ingest for now."],
        ["doc_type", "manual | guide | toolkit | report | course | other"],
        ["tags", "Comma-separated topics, e.g. 'icebreakers, teambuilding'. Used in the chunk header."],
        ["summary / who_it_helps / how_it_helps", "Doc-level context. For non-EN docs, write an EN summary."],
        ["source_link", "Where to find the original (shown in citations)."],
        ["metadata_source", "How this row was prefilled. 'MANUAL - fill me' / 'fuzzy?' = check it."],
    ]
    for row in notes:
        rm.append(row)
    rm.column_dimensions["A"].width = 40
    rm.column_dimensions["B"].width = 90
    for cell in rm[1]:
        cell.font = Font(bold=True)

    wb.save(out_path)


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("--partners", default=os.path.join(BASE_DIR, "data", "metadata", "Partners_data_assets.xlsx"))
    ap.add_argument("--cazalla", default=os.path.join(BASE_DIR, "data", "metadata", "Cazala_sources_metadata.xlsx"))
    ap.add_argument("--out", default=DEFAULT_OUT)
    build(ap.parse_args())
