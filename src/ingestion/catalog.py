"""
catalog.py
==========
The document-level metadata catalog is the human-editable source of truth for
per-PDF metadata. build_catalog.py writes/updates it; ingest_pdfs.py reads it
and stamps every chunk with its document's row.

One row per PDF. Columns below. Keep it in data/metadata/docs_catalog.xlsx.
"""

import re
import unicodedata

CATALOG_COLUMNS = [
    "file_name",        # exact PDF filename in data/raw_pdfs/ (the join key)
    "title",            # human-readable title
    "language",         # ISO code: en, es, de, fr, it, uk, ...
    "country",
    "doc_type",         # manual | guide | toolkit | report | course | other
    "tags",             # comma-separated, e.g. "icebreakers, teambuilding"
    "summary",          # short description (for non-EN docs, an EN summary)
    "who_it_helps",
    "how_it_helps",
    "source_link",
    "year",
    "metadata_source",  # partners | cazalla | fuzzy? | MANUAL - fill me
]

# Languages we actually ingest in Step 1 (all-MiniLM is English-centric).
# Others (uk, bg, ru, el, ...) are skipped for now; revisit with bge-m3 +
# summary-based ingestion. See the Ukrainian discussion.
SUPPORTED_LANGS = {"en", "es", "de", "fr", "it", "pt", "nl"}


def normalize(s: str) -> str:
    """Lowercase, strip accents & non-alphanumerics — for fuzzy matching."""
    if not s:
        return ""
    s = unicodedata.normalize("NFKD", str(s))
    s = "".join(c for c in s if not unicodedata.combining(c))
    return re.sub(r"[^a-z0-9]+", "", s.lower())


def load_catalog(path):
    """Return {file_name: {column: value}} from docs_catalog.xlsx, or {} if absent."""
    import os
    if not os.path.exists(path):
        return {}
    import openpyxl
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb["catalog"] if "catalog" in wb.sheetnames else wb.worksheets[0]
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return {}
    header = [str(c).strip() if c is not None else "" for c in rows[0]]
    out = {}
    for r in rows[1:]:
        row = {header[i]: (r[i] if i < len(r) else None) for i in range(len(header))}
        fn = (row.get("file_name") or "").strip()
        if fn:
            out[fn] = row
    return out
